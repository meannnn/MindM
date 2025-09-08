"""
图片解析器

专门用于解析Word文档中的图片和形状信息。
"""

import logging
import os
from typing import Dict, List, Any, Optional
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)


class ImageParser:
    """图片解析器"""
    
    def __init__(self, document: Document):
        """
        初始化图片解析器
        
        Args:
            document: Word文档对象
        """
        self.document = document
        self.images = []
        self.shapes = []
    
    def parse_all_images(self) -> List[Dict[str, Any]]:
        """
        解析所有图片
        
        Returns:
            List[Dict[str, Any]]: 所有图片的解析数据
        """
        images_data = []
        
        try:
            # 解析内联形状（图片）
            inline_images = self._parse_inline_shapes()
            images_data.extend(inline_images)
            
            # 解析段落中的图片
            paragraph_images = self._parse_paragraph_images()
            images_data.extend(paragraph_images)
            
            # 解析形状中的图片
            shape_images = self._parse_shape_images()
            images_data.extend(shape_images)
            
            self.images = images_data
            logger.info(f"成功解析 {len(images_data)} 个图片")
            return images_data
            
        except Exception as e:
            logger.error(f"解析图片失败: {e}")
            return []
    
    def _parse_inline_shapes(self) -> List[Dict[str, Any]]:
        """
        解析内联形状（图片）
        
        Returns:
            List[Dict[str, Any]]: 内联图片数据
        """
        inline_images = []
        
        try:
            for i, shape in enumerate(self.document.inline_shapes):
                if hasattr(shape, 'image'):
                    image_data = {
                        'index': i,
                        'type': 'inline',
                        'width': shape.width,
                        'height': shape.height,
                        'width_inches': shape.width.inches,
                        'height_inches': shape.height.inches,
                        'filename': getattr(shape.image, 'filename', 'unknown'),
                        'content_type': getattr(shape.image, 'content_type', 'unknown'),
                        'blip_id': self._get_blip_id(shape),
                        'position': self._get_image_position(shape)
                    }
                    inline_images.append(image_data)
            
            logger.info(f"解析了 {len(inline_images)} 个内联图片")
            return inline_images
            
        except Exception as e:
            logger.error(f"解析内联形状失败: {e}")
            return []
    
    def _parse_paragraph_images(self) -> List[Dict[str, Any]]:
        """
        解析段落中的图片
        
        Returns:
            List[Dict[str, Any]]: 段落图片数据
        """
        paragraph_images = []
        
        try:
            for para_idx, paragraph in enumerate(self.document.paragraphs):
                for run_idx, run in enumerate(paragraph.runs):
                    # 检查运行中是否包含图片
                    if self._has_image_in_run(run):
                        image_data = {
                            'index': len(paragraph_images),
                            'type': 'paragraph',
                            'paragraph_index': para_idx,
                            'run_index': run_idx,
                            'paragraph_text': paragraph.text.strip(),
                            'run_text': run.text.strip(),
                            'position': f"段落{para_idx}, 运行{run_idx}"
                        }
                        paragraph_images.append(image_data)
            
            logger.info(f"解析了 {len(paragraph_images)} 个段落图片")
            return paragraph_images
            
        except Exception as e:
            logger.error(f"解析段落图片失败: {e}")
            return []
    
    def _parse_shape_images(self) -> List[Dict[str, Any]]:
        """
        解析形状中的图片
        
        Returns:
            List[Dict[str, Any]]: 形状图片数据
        """
        shape_images = []
        
        try:
            # 注意：python-docx对形状的支持有限
            # 这里主要处理一些基本的形状信息
            for para_idx, paragraph in enumerate(self.document.paragraphs):
                # 检查段落中是否有形状相关的XML
                if self._has_shapes_in_paragraph(paragraph):
                    shape_data = {
                        'index': len(shape_images),
                        'type': 'shape',
                        'paragraph_index': para_idx,
                        'paragraph_text': paragraph.text.strip(),
                        'position': f"段落{para_idx}"
                    }
                    shape_images.append(shape_data)
            
            logger.info(f"解析了 {len(shape_images)} 个形状图片")
            return shape_images
            
        except Exception as e:
            logger.error(f"解析形状图片失败: {e}")
            return []
    
    def _get_blip_id(self, shape) -> Optional[str]:
        """
        获取图片的blip ID
        
        Args:
            shape: 形状对象
            
        Returns:
            Optional[str]: blip ID
        """
        try:
            if hasattr(shape, '_inline'):
                inline = shape._inline
                if hasattr(inline, 'graphic'):
                    graphic = inline.graphic
                    if hasattr(graphic, 'graphicData'):
                        graphic_data = graphic.graphicData
                        # 查找blip元素
                        blip_elements = graphic_data.xpath('.//a:blip', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
                        if blip_elements:
                            return blip_elements[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            
            return None
            
        except Exception as e:
            logger.error(f"获取blip ID失败: {e}")
            return None
    
    def _get_image_position(self, shape) -> Dict[str, Any]:
        """
        获取图片位置信息
        
        Args:
            shape: 形状对象
            
        Returns:
            Dict[str, Any]: 位置信息
        """
        try:
            position = {
                'anchor': None,
                'wrap': None,
                'z_order': None
            }
            
            if hasattr(shape, '_inline'):
                inline = shape._inline
                if hasattr(inline, 'anchor'):
                    position['anchor'] = 'inline'
                elif hasattr(inline, 'wrap'):
                    position['wrap'] = str(inline.wrap)
            
            return position
            
        except Exception as e:
            logger.error(f"获取图片位置失败: {e}")
            return {}
    
    def _has_image_in_run(self, run) -> bool:
        """
        检查运行中是否包含图片
        
        Args:
            run: 文本运行对象
            
        Returns:
            bool: 是否包含图片
        """
        try:
            # 检查XML中是否包含图片相关元素
            run_xml = run._element.xml
            return ('a:blip' in run_xml or 
                   'pic:pic' in run_xml or 
                   'w:drawing' in run_xml)
            
        except Exception as e:
            logger.error(f"检查运行图片失败: {e}")
            return False
    
    def _has_shapes_in_paragraph(self, paragraph) -> bool:
        """
        检查段落中是否包含形状
        
        Args:
            paragraph: 段落对象
            
        Returns:
            bool: 是否包含形状
        """
        try:
            # 检查段落XML中是否包含形状相关元素
            para_xml = paragraph._element.xml
            return ('w:drawing' in para_xml or 
                   'w:object' in para_xml or
                   'w:control' in para_xml)
            
        except Exception as e:
            logger.error(f"检查段落形状失败: {e}")
            return False
    
    def extract_image_info(self, image_index: int) -> Dict[str, Any]:
        """
        提取指定图片的详细信息
        
        Args:
            image_index: 图片索引
            
        Returns:
            Dict[str, Any]: 图片详细信息
        """
        try:
            if image_index >= len(self.images):
                return {}
            
            image_data = self.images[image_index]
            
            # 添加更多详细信息
            detailed_info = image_data.copy()
            detailed_info.update({
                'aspect_ratio': self._calculate_aspect_ratio(image_data),
                'size_category': self._categorize_size(image_data),
                'orientation': self._determine_orientation(image_data),
                'estimated_pixels': self._estimate_pixels(image_data)
            })
            
            return detailed_info
            
        except Exception as e:
            logger.error(f"提取图片信息失败: {e}")
            return {}
    
    def _calculate_aspect_ratio(self, image_data: Dict[str, Any]) -> float:
        """
        计算图片宽高比
        
        Args:
            image_data: 图片数据
            
        Returns:
            float: 宽高比
        """
        try:
            if 'width_inches' in image_data and 'height_inches' in image_data:
                width = image_data['width_inches']
                height = image_data['height_inches']
                if height > 0:
                    return round(width / height, 2)
            
            return 1.0
            
        except Exception as e:
            logger.error(f"计算宽高比失败: {e}")
            return 1.0
    
    def _categorize_size(self, image_data: Dict[str, Any]) -> str:
        """
        分类图片大小
        
        Args:
            image_data: 图片数据
            
        Returns:
            str: 大小分类
        """
        try:
            if 'width_inches' in image_data and 'height_inches' in image_data:
                width = image_data['width_inches']
                height = image_data['height_inches']
                area = width * height
                
                if area < 1:
                    return 'small'
                elif area < 4:
                    return 'medium'
                elif area < 9:
                    return 'large'
                else:
                    return 'extra_large'
            
            return 'unknown'
            
        except Exception as e:
            logger.error(f"分类图片大小失败: {e}")
            return 'unknown'
    
    def _determine_orientation(self, image_data: Dict[str, Any]) -> str:
        """
        确定图片方向
        
        Args:
            image_data: 图片数据
            
        Returns:
            str: 方向（portrait/landscape/square）
        """
        try:
            if 'width_inches' in image_data and 'height_inches' in image_data:
                width = image_data['width_inches']
                height = image_data['height_inches']
                
                if abs(width - height) < 0.1:
                    return 'square'
                elif width > height:
                    return 'landscape'
                else:
                    return 'portrait'
            
            return 'unknown'
            
        except Exception as e:
            logger.error(f"确定图片方向失败: {e}")
            return 'unknown'
    
    def _estimate_pixels(self, image_data: Dict[str, Any]) -> Dict[str, int]:
        """
        估算图片像素尺寸
        
        Args:
            image_data: 图片数据
            
        Returns:
            Dict[str, int]: 像素尺寸
        """
        try:
            if 'width_inches' in image_data and 'height_inches' in image_data:
                width_inches = image_data['width_inches']
                height_inches = image_data['height_inches']
                
                # 假设标准DPI为96
                dpi = 96
                width_pixels = int(width_inches * dpi)
                height_pixels = int(height_inches * dpi)
                
                return {
                    'width': width_pixels,
                    'height': height_pixels,
                    'total_pixels': width_pixels * height_pixels
                }
            
            return {'width': 0, 'height': 0, 'total_pixels': 0}
            
        except Exception as e:
            logger.error(f"估算像素尺寸失败: {e}")
            return {'width': 0, 'height': 0, 'total_pixels': 0}
    
    def get_images_summary(self) -> Dict[str, Any]:
        """
        获取图片汇总信息
        
        Returns:
            Dict[str, Any]: 图片汇总信息
        """
        try:
            if not self.images:
                self.parse_all_images()
            
            summary = {
                'total_count': len(self.images),
                'by_type': {},
                'by_size': {},
                'by_orientation': {},
                'total_area': 0,
                'average_size': 0
            }
            
            # 按类型统计
            for image in self.images:
                img_type = image.get('type', 'unknown')
                summary['by_type'][img_type] = summary['by_type'].get(img_type, 0) + 1
            
            # 按大小统计
            for image in self.images:
                size_cat = self._categorize_size(image)
                summary['by_size'][size_cat] = summary['by_size'].get(size_cat, 0) + 1
            
            # 按方向统计
            for image in self.images:
                orientation = self._determine_orientation(image)
                summary['by_orientation'][orientation] = summary['by_orientation'].get(orientation, 0) + 1
            
            # 计算总面积
            total_area = 0
            for image in self.images:
                if 'width_inches' in image and 'height_inches' in image:
                    total_area += image['width_inches'] * image['height_inches']
            
            summary['total_area'] = round(total_area, 2)
            summary['average_size'] = round(total_area / len(self.images), 2) if self.images else 0
            
            return summary
            
        except Exception as e:
            logger.error(f"获取图片汇总失败: {e}")
            return {}
    
    def export_images_to_list(self) -> List[Dict[str, Any]]:
        """
        导出所有图片信息为列表格式
        
        Returns:
            List[Dict[str, Any]]: 图片信息列表
        """
        try:
            if not self.images:
                self.parse_all_images()
            
            export_list = []
            for i, image in enumerate(self.images):
                image_info = self.extract_image_info(i)
                export_list.append(image_info)
            
            return export_list
            
        except Exception as e:
            logger.error(f"导出图片列表失败: {e}")
            return []


# 使用示例
if __name__ == "__main__":
    # 配置日志
    logging.basicConfig(level=logging.INFO)
    
    # 加载文档
    doc = Document("example.docx")
    
    # 创建图片解析器
    parser = ImageParser(doc)
    
    # 解析所有图片
    images_data = parser.parse_all_images()
    
    # 打印图片信息
    print(f"文档中共有 {len(images_data)} 个图片")
    
    for i, image in enumerate(images_data):
        print(f"图片 {i}:")
        print(f"  类型: {image['type']}")
        if 'width_inches' in image:
            print(f"  尺寸: {image['width_inches']:.2f} x {image['height_inches']:.2f} 英寸")
        print(f"  位置: {image.get('position', '未知')}")
        print()
    
    # 获取汇总信息
    summary = parser.get_images_summary()
    print("图片汇总:")
    print(f"  总数: {summary['total_count']}")
    print(f"  按类型: {summary['by_type']}")
    print(f"  按大小: {summary['by_size']}")
    print(f"  按方向: {summary['by_orientation']}")
    print(f"  总面积: {summary['total_area']} 平方英寸")