"""
表格解析器

专门用于解析Word文档中的表格数据，提供详细的表格分析功能。
"""

import logging
from typing import Dict, List, Any, Optional
from docx import Document
from docx.table import Table, _Cell, _Row
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

logger = logging.getLogger(__name__)


class TableParser:
    """表格解析器"""
    
    def __init__(self, document: Document):
        """
        初始化表格解析器
        
        Args:
            document: Word文档对象
        """
        self.document = document
        self.tables = document.tables
    
    def parse_all_tables(self) -> List[Dict[str, Any]]:
        """
        解析所有表格
        
        Returns:
            List[Dict[str, Any]]: 所有表格的解析数据
        """
        tables_data = []
        
        try:
            for i, table in enumerate(self.tables):
                table_data = self.parse_table(table, i)
                tables_data.append(table_data)
            
            logger.info(f"成功解析 {len(tables_data)} 个表格")
            return tables_data
            
        except Exception as e:
            logger.error(f"解析表格失败: {e}")
            return []
    
    def parse_table(self, table: Table, index: int = 0) -> Dict[str, Any]:
        """
        解析单个表格
        
        Args:
            table: 表格对象
            index: 表格索引
            
        Returns:
            Dict[str, Any]: 表格解析数据
        """
        try:
            table_data = {
                'index': index,
                'rows': len(table.rows),
                'columns': len(table.columns),
                'data': self._extract_table_data(table),
                'structure': self._analyze_table_structure(table),
                'style': self._get_table_style(table),
                'alignment': self._get_table_alignment(table),
                'headers': self._identify_headers(table),
                'merged_cells': self._find_merged_cells(table)
            }
            
            return table_data
            
        except Exception as e:
            logger.error(f"解析表格 {index} 失败: {e}")
            return {}
    
    def _extract_table_data(self, table: Table) -> List[List[Dict[str, Any]]]:
        """
        提取表格数据
        
        Args:
            table: 表格对象
            
        Returns:
            List[List[Dict[str, Any]]]: 表格数据矩阵
        """
        table_data = []
        
        try:
            for row_idx, row in enumerate(table.rows):
                row_data = []
                for col_idx, cell in enumerate(row.cells):
                    cell_data = {
                        'row': row_idx,
                        'column': col_idx,
                        'text': cell.text.strip(),
                        'paragraphs': [p.text.strip() for p in cell.paragraphs if p.text.strip()],
                        'runs': self._extract_cell_runs(cell),
                        'formatting': self._get_cell_formatting(cell)
                    }
                    row_data.append(cell_data)
                table_data.append(row_data)
            
            return table_data
            
        except Exception as e:
            logger.error(f"提取表格数据失败: {e}")
            return []
    
    def _extract_cell_runs(self, cell: _Cell) -> List[Dict[str, Any]]:
        """
        提取单元格中的文本运行信息
        
        Args:
            cell: 单元格对象
            
        Returns:
            List[Dict[str, Any]]: 文本运行信息
        """
        runs_data = []
        
        try:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        run_data = {
                            'text': run.text,
                            'bold': run.bold,
                            'italic': run.italic,
                            'underline': run.underline,
                            'font_name': run.font.name,
                            'font_size': run.font.size.pt if run.font.size else None,
                            'font_color': str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None
                        }
                        runs_data.append(run_data)
            
            return runs_data
            
        except Exception as e:
            logger.error(f"提取单元格文本运行失败: {e}")
            return []
    
    def _get_cell_formatting(self, cell: _Cell) -> Dict[str, Any]:
        """
        获取单元格格式信息
        
        Args:
            cell: 单元格对象
            
        Returns:
            Dict[str, Any]: 格式信息
        """
        try:
            formatting = {
                'vertical_alignment': self._get_vertical_alignment_name(cell.vertical_alignment),
                'shading': self._get_cell_shading(cell),
                'borders': self._get_cell_borders(cell),
                'margins': self._get_cell_margins(cell)
            }
            
            return formatting
            
        except Exception as e:
            logger.error(f"获取单元格格式失败: {e}")
            return {}
    
    def _get_vertical_alignment_name(self, alignment) -> str:
        """获取垂直对齐方式名称"""
        if alignment == WD_CELL_VERTICAL_ALIGNMENT.TOP:
            return 'top'
        elif alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER:
            return 'center'
        elif alignment == WD_CELL_VERTICAL_ALIGNMENT.BOTTOM:
            return 'bottom'
        else:
            return 'top'
    
    def _get_cell_shading(self, cell: _Cell) -> Dict[str, Any]:
        """获取单元格底纹信息"""
        try:
            shading = cell._tc.get_or_add_tcPr().get_or_add_shd()
            return {
                'color': str(shading.val) if shading.val else None,
                'fill': str(shading.fill) if shading.fill else None
            }
        except:
            return {}
    
    def _get_cell_borders(self, cell: _Cell) -> Dict[str, Any]:
        """获取单元格边框信息"""
        try:
            borders = {}
            tc_borders = cell._tc.get_or_add_tcPr().get_or_add_tcBorders()
            
            for border_name in ['top', 'bottom', 'left', 'right']:
                border = getattr(tc_borders, f'get_or_add_{border_name}')()
                borders[border_name] = {
                    'color': str(border.color) if border.color else None,
                    'size': str(border.sz) if border.sz else None,
                    'style': str(border.val) if border.val else None
                }
            
            return borders
        except:
            return {}
    
    def _get_cell_margins(self, cell: _Cell) -> Dict[str, Any]:
        """获取单元格边距信息"""
        try:
            margins = {}
            tc_mar = cell._tc.get_or_add_tcPr().get_or_add_tcMar()
            
            for margin_name in ['top', 'bottom', 'left', 'right']:
                margin = getattr(tc_mar, f'get_or_add_{margin_name}')()
                margins[margin_name] = str(margin.w) if margin.w else '0'
            
            return margins
        except:
            return {}
    
    def _analyze_table_structure(self, table: Table) -> Dict[str, Any]:
        """
        分析表格结构
        
        Args:
            table: 表格对象
            
        Returns:
            Dict[str, Any]: 表格结构信息
        """
        try:
            structure = {
                'has_header': False,
                'header_rows': 0,
                'data_rows': 0,
                'empty_cells': 0,
                'merged_cells': 0,
                'column_types': self._analyze_column_types(table),
                'row_types': self._analyze_row_types(table)
            }
            
            # 分析表头
            if len(table.rows) > 0:
                first_row = table.rows[0]
                if self._is_header_row(first_row):
                    structure['has_header'] = True
                    structure['header_rows'] = 1
            
            # 统计空单元格
            for row in table.rows:
                for cell in row.cells:
                    if not cell.text.strip():
                        structure['empty_cells'] += 1
            
            structure['data_rows'] = len(table.rows) - structure['header_rows']
            
            return structure
            
        except Exception as e:
            logger.error(f"分析表格结构失败: {e}")
            return {}
    
    def _is_header_row(self, row: _Row) -> bool:
        """
        判断是否为表头行
        
        Args:
            row: 行对象
            
        Returns:
            bool: 是否为表头行
        """
        try:
            # 检查是否有粗体文本
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.bold and run.text.strip():
                            return True
            
            # 检查是否包含常见的表头关键词
            header_keywords = ['序号', '编号', '名称', '标题', '类型', '状态', '时间', '日期', '数量', '金额']
            row_text = ' '.join([cell.text.strip() for cell in row.cells])
            
            for keyword in header_keywords:
                if keyword in row_text:
                    return True
            
            return False
            
        except:
            return False
    
    def _analyze_column_types(self, table: Table) -> List[str]:
        """
        分析列类型
        
        Args:
            table: 表格对象
            
        Returns:
            List[str]: 每列的类型
        """
        column_types = []
        
        try:
            if len(table.rows) == 0:
                return column_types
            
            num_cols = len(table.columns)
            
            for col_idx in range(num_cols):
                column_data = []
                for row in table.rows:
                    if col_idx < len(row.cells):
                        cell_text = row.cells[col_idx].text.strip()
                        if cell_text:
                            column_data.append(cell_text)
                
                # 分析列类型
                col_type = self._determine_column_type(column_data)
                column_types.append(col_type)
            
            return column_types
            
        except Exception as e:
            logger.error(f"分析列类型失败: {e}")
            return []
    
    def _analyze_row_types(self, table: Table) -> List[str]:
        """
        分析行类型
        
        Args:
            table: 表格对象
            
        Returns:
            List[str]: 每行的类型
        """
        row_types = []
        
        try:
            for row_idx, row in enumerate(table.rows):
                if self._is_header_row(row):
                    row_types.append('header')
                else:
                    row_types.append('data')
            
            return row_types
            
        except Exception as e:
            logger.error(f"分析行类型失败: {e}")
            return []
    
    def _determine_column_type(self, column_data: List[str]) -> str:
        """
        确定列的数据类型
        
        Args:
            column_data: 列数据
            
        Returns:
            str: 列类型
        """
        if not column_data:
            return 'empty'
        
        # 检查是否为数字
        numeric_count = 0
        for item in column_data:
            try:
                float(item.replace(',', '').replace('%', ''))
                numeric_count += 1
            except:
                pass
        
        if numeric_count / len(column_data) > 0.8:
            return 'numeric'
        
        # 检查是否为日期
        date_count = 0
        import re
        date_pattern = r'\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4}'
        for item in column_data:
            if re.match(date_pattern, item):
                date_count += 1
        
        if date_count / len(column_data) > 0.5:
            return 'date'
        
        # 检查是否为布尔值
        boolean_count = 0
        boolean_values = ['是', '否', 'true', 'false', 'yes', 'no', 'y', 'n', '1', '0']
        for item in column_data:
            if item.lower() in boolean_values:
                boolean_count += 1
        
        if boolean_count / len(column_data) > 0.8:
            return 'boolean'
        
        return 'text'
    
    def _get_table_style(self, table: Table) -> Dict[str, Any]:
        """
        获取表格样式信息
        
        Args:
            table: 表格对象
            
        Returns:
            Dict[str, Any]: 样式信息
        """
        try:
            style_info = {
                'name': table.style.name if table.style else 'Table Grid',
                'builtin': table.style.builtin if table.style else True,
                'type': str(table.style.type) if table.style else 'table'
            }
            
            return style_info
            
        except Exception as e:
            logger.error(f"获取表格样式失败: {e}")
            return {}
    
    def _get_table_alignment(self, table: Table) -> str:
        """
        获取表格对齐方式
        
        Args:
            table: 表格对象
            
        Returns:
            str: 对齐方式
        """
        if table.alignment == WD_TABLE_ALIGNMENT.LEFT:
            return 'left'
        elif table.alignment == WD_TABLE_ALIGNMENT.CENTER:
            return 'center'
        elif table.alignment == WD_TABLE_ALIGNMENT.RIGHT:
            return 'right'
        else:
            return 'left'
    
    def _identify_headers(self, table: Table) -> List[Dict[str, Any]]:
        """
        识别表格表头
        
        Args:
            table: 表格对象
            
        Returns:
            List[Dict[str, Any]]: 表头信息
        """
        headers = []
        
        try:
            if len(table.rows) > 0:
                header_row = table.rows[0]
                for col_idx, cell in enumerate(header_row.cells):
                    header_info = {
                        'column': col_idx,
                        'text': cell.text.strip(),
                        'is_header': True
                    }
                    headers.append(header_info)
            
            return headers
            
        except Exception as e:
            logger.error(f"识别表头失败: {e}")
            return []
    
    def _find_merged_cells(self, table: Table) -> List[Dict[str, Any]]:
        """
        查找合并单元格
        
        Args:
            table: 表格对象
            
        Returns:
            List[Dict[str, Any]]: 合并单元格信息
        """
        merged_cells = []
        
        try:
            for row_idx, row in enumerate(table.rows):
                for col_idx, cell in enumerate(row.cells):
                    # 检查是否为合并单元格
                    if hasattr(cell, '_tc'):
                        tc = cell._tc
                        if hasattr(tc, 'grid_span') and tc.grid_span > 1:
                            merged_info = {
                                'row': row_idx,
                                'column': col_idx,
                                'span': tc.grid_span,
                                'text': cell.text.strip()
                            }
                            merged_cells.append(merged_info)
            
            return merged_cells
            
        except Exception as e:
            logger.error(f"查找合并单元格失败: {e}")
            return []
    
    def export_table_to_dict(self, table_index: int) -> Dict[str, Any]:
        """
        将表格导出为字典格式
        
        Args:
            table_index: 表格索引
            
        Returns:
            Dict[str, Any]: 表格字典数据
        """
        try:
            if table_index >= len(self.tables):
                return {}
            
            table = self.tables[table_index]
            table_data = self.parse_table(table, table_index)
            
            # 转换为更简洁的字典格式
            export_data = {
                'headers': [cell['text'] for cell in table_data['headers']],
                'rows': []
            }
            
            # 跳过表头行，只导出数据行
            start_row = 1 if table_data['structure']['has_header'] else 0
            
            for row_data in table_data['data'][start_row:]:
                row_dict = {}
                for i, cell_data in enumerate(row_data):
                    if i < len(export_data['headers']):
                        row_dict[export_data['headers'][i]] = cell_data['text']
                export_data['rows'].append(row_dict)
            
            return export_data
            
        except Exception as e:
            logger.error(f"导出表格失败: {e}")
            return {}
    
    def export_all_tables_to_dict(self) -> List[Dict[str, Any]]:
        """
        导出所有表格为字典格式
        
        Returns:
            List[Dict[str, Any]]: 所有表格的字典数据
        """
        all_tables = []
        
        try:
            for i in range(len(self.tables)):
                table_dict = self.export_table_to_dict(i)
                if table_dict:
                    all_tables.append(table_dict)
            
            return all_tables
            
        except Exception as e:
            logger.error(f"导出所有表格失败: {e}")
            return []


# 使用示例
if __name__ == "__main__":
    # 配置日志
    logging.basicConfig(level=logging.INFO)
    
    # 加载文档
    doc = Document("example.docx")
    
    # 创建表格解析器
    parser = TableParser(doc)
    
    # 解析所有表格
    tables_data = parser.parse_all_tables()
    
    # 打印表格信息
    for i, table_data in enumerate(tables_data):
        print(f"表格 {i}:")
        print(f"  行数: {table_data['rows']}")
        print(f"  列数: {table_data['columns']}")
        print(f"  是否有表头: {table_data['structure']['has_header']}")
        print(f"  列类型: {table_data['structure']['column_types']}")
        print()
    
    # 导出为字典格式
    tables_dict = parser.export_all_tables_to_dict()
    print(f"导出了 {len(tables_dict)} 个表格")