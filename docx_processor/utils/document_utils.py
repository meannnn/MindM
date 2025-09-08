"""
文档工具函数

提供文档处理相关的实用工具函数。
"""

import os
import logging
from typing import Dict, List, Any, Optional, Union
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

logger = logging.getLogger(__name__)


class DocumentUtils:
    """文档工具类"""
    
    @staticmethod
    def validate_docx_file(file_path: str) -> Dict[str, Any]:
        """
        验证docx文件
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict[str, Any]: 验证结果
        """
        result = {
            'is_valid': False,
            'errors': [],
            'warnings': [],
            'file_info': {}
        }
        
        try:
            # 检查文件是否存在
            if not os.path.exists(file_path):
                result['errors'].append("文件不存在")
                return result
            
            # 检查文件扩展名
            if not file_path.lower().endswith('.docx'):
                result['errors'].append("文件不是.docx格式")
                return result
            
            # 检查文件大小
            file_size = os.path.getsize(file_path)
            if file_size == 0:
                result['errors'].append("文件为空")
                return result
            
            if file_size > 50 * 1024 * 1024:  # 50MB
                result['warnings'].append("文件过大，可能影响处理性能")
            
            # 尝试打开文档
            try:
                doc = Document(file_path)
                result['file_info'] = {
                    'size': file_size,
                    'paragraph_count': len(doc.paragraphs),
                    'table_count': len(doc.tables),
                    'section_count': len(doc.sections)
                }
                result['is_valid'] = True
                
            except Exception as e:
                result['errors'].append(f"无法打开文档: {str(e)}")
            
        except Exception as e:
            result['errors'].append(f"验证文件时发生错误: {str(e)}")
        
        return result
    
    @staticmethod
    def get_document_summary(file_path: str) -> Dict[str, Any]:
        """
        获取文档摘要信息
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict[str, Any]: 文档摘要
        """
        summary = {
            'file_path': file_path,
            'file_name': os.path.basename(file_path),
            'file_size': 0,
            'paragraph_count': 0,
            'table_count': 0,
            'image_count': 0,
            'word_count': 0,
            'character_count': 0,
            'title': '',
            'author': '',
            'created': None,
            'modified': None
        }
        
        try:
            # 文件基本信息
            if os.path.exists(file_path):
                summary['file_size'] = os.path.getsize(file_path)
            
            # 文档内容信息
            doc = Document(file_path)
            
            # 段落和文本统计
            summary['paragraph_count'] = len(doc.paragraphs)
            
            # 统计单词和字符数
            word_count = 0
            char_count = 0
            
            for paragraph in doc.paragraphs:
                text = paragraph.text
                char_count += len(text)
                word_count += len(text.split())
            
            summary['word_count'] = word_count
            summary['character_count'] = char_count
            
            # 表格统计
            summary['table_count'] = len(doc.tables)
            
            # 图片统计
            summary['image_count'] = len(doc.inline_shapes)
            
            # 文档属性
            core_props = doc.core_properties
            summary['title'] = core_props.title or ''
            summary['author'] = core_props.author or ''
            summary['created'] = core_props.created
            summary['modified'] = core_props.modified
            
        except Exception as e:
            logger.error(f"获取文档摘要失败: {e}")
        
        return summary
    
    @staticmethod
    def extract_text_content(file_path: str, include_tables: bool = True) -> str:
        """
        提取文档的纯文本内容
        
        Args:
            file_path: 文件路径
            include_tables: 是否包含表格内容
            
        Returns:
            str: 纯文本内容
        """
        try:
            doc = Document(file_path)
            text_content = []
            
            # 提取段落文本
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            # 提取表格文本
            if include_tables:
                for table in doc.tables:
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_text.append(cell.text.strip())
                        if row_text:
                            text_content.append('\t'.join(row_text))
            
            return '\n'.join(text_content)
            
        except Exception as e:
            logger.error(f"提取文本内容失败: {e}")
            return ""
    
    @staticmethod
    def extract_structured_content(file_path: str) -> Dict[str, Any]:
        """
        提取结构化内容
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict[str, Any]: 结构化内容
        """
        try:
            doc = Document(file_path)
            structured_content = {
                'headings': [],
                'paragraphs': [],
                'tables': [],
                'lists': [],
                'images': []
            }
            
            # 提取标题
            for paragraph in doc.paragraphs:
                if paragraph.style and paragraph.style.name.startswith('Heading'):
                    structured_content['headings'].append({
                        'text': paragraph.text,
                        'level': int(paragraph.style.name.split()[-1]) if paragraph.style.name.split()[-1].isdigit() else 1,
                        'style': paragraph.style.name
                    })
            
            # 提取段落
            for paragraph in doc.paragraphs:
                if paragraph.text.strip() and not (paragraph.style and paragraph.style.name.startswith('Heading')):
                    structured_content['paragraphs'].append({
                        'text': paragraph.text,
                        'style': paragraph.style.name if paragraph.style else 'Normal'
                    })
            
            # 提取表格
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                structured_content['tables'].append(table_data)
            
            # 提取列表
            for paragraph in doc.paragraphs:
                if paragraph.style and ('List' in paragraph.style.name or 'Bullet' in paragraph.style.name):
                    structured_content['lists'].append({
                        'text': paragraph.text,
                        'style': paragraph.style.name
                    })
            
            # 提取图片信息
            for i, shape in enumerate(doc.inline_shapes):
                structured_content['images'].append({
                    'index': i,
                    'width': shape.width,
                    'height': shape.height
                })
            
            return structured_content
            
        except Exception as e:
            logger.error(f"提取结构化内容失败: {e}")
            return {}
    
    @staticmethod
    def compare_documents(doc1_path: str, doc2_path: str) -> Dict[str, Any]:
        """
        比较两个文档
        
        Args:
            doc1_path: 第一个文档路径
            doc2_path: 第二个文档路径
            
        Returns:
            Dict[str, Any]: 比较结果
        """
        try:
            # 获取两个文档的摘要
            summary1 = DocumentUtils.get_document_summary(doc1_path)
            summary2 = DocumentUtils.get_document_summary(doc2_path)
            
            # 比较结果
            comparison = {
                'doc1': summary1,
                'doc2': summary2,
                'differences': {
                    'paragraph_count': summary1['paragraph_count'] - summary2['paragraph_count'],
                    'table_count': summary1['table_count'] - summary2['table_count'],
                    'word_count': summary1['word_count'] - summary2['word_count'],
                    'character_count': summary1['character_count'] - summary2['character_count']
                },
                'similarities': {
                    'same_author': summary1['author'] == summary2['author'],
                    'same_title': summary1['title'] == summary2['title']
                }
            }
            
            return comparison
            
        except Exception as e:
            logger.error(f"比较文档失败: {e}")
            return {}
    
    @staticmethod
    def merge_documents(doc_paths: List[str], output_path: str) -> bool:
        """
        合并多个文档
        
        Args:
            doc_paths: 文档路径列表
            output_path: 输出文件路径
            
        Returns:
            bool: 是否合并成功
        """
        try:
            if not doc_paths:
                logger.error("没有提供要合并的文档")
                return False
            
            # 创建新文档
            merged_doc = Document()
            
            for i, doc_path in enumerate(doc_paths):
                if not os.path.exists(doc_path):
                    logger.warning(f"文档不存在，跳过: {doc_path}")
                    continue
                
                # 加载文档
                doc = Document(doc_path)
                
                # 添加文档标题
                if i > 0:
                    merged_doc.add_page_break()
                
                merged_doc.add_heading(f"文档 {i+1}: {os.path.basename(doc_path)}", level=1)
                
                # 复制段落
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        new_para = merged_doc.add_paragraph(paragraph.text)
                        if paragraph.style:
                            new_para.style = paragraph.style.name
                
                # 复制表格
                for table in doc.tables:
                    new_table = merged_doc.add_table(rows=len(table.rows), cols=len(table.columns))
                    for i, row in enumerate(table.rows):
                        for j, cell in enumerate(row.cells):
                            if i < len(new_table.rows) and j < len(new_table.rows[i].cells):
                                new_table.rows[i].cells[j].text = cell.text
            
            # 保存合并后的文档
            merged_doc.save(output_path)
            logger.info(f"成功合并 {len(doc_paths)} 个文档到: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"合并文档失败: {e}")
            return False
    
    @staticmethod
    def split_document_by_headings(file_path: str, output_dir: str) -> List[str]:
        """
        按标题分割文档
        
        Args:
            file_path: 源文档路径
            output_dir: 输出目录
            
        Returns:
            List[str]: 生成的文件路径列表
        """
        try:
            doc = Document(file_path)
            output_files = []
            
            # 确保输出目录存在
            if not os.path.exists(output_dir):
                os.makedirs(output_dir)
            
            current_doc = None
            current_heading = None
            file_counter = 1
            
            for paragraph in doc.paragraphs:
                # 检查是否为标题
                if paragraph.style and paragraph.style.name.startswith('Heading'):
                    # 保存当前文档
                    if current_doc:
                        output_file = os.path.join(output_dir, f"section_{file_counter}_{current_heading}.docx")
                        current_doc.save(output_file)
                        output_files.append(output_file)
                        file_counter += 1
                    
                    # 创建新文档
                    current_doc = Document()
                    current_heading = paragraph.text.replace('/', '_').replace('\\', '_')[:50]
                    current_doc.add_heading(paragraph.text, level=1)
                
                # 添加段落到当前文档
                elif current_doc:
                    current_doc.add_paragraph(paragraph.text)
            
            # 保存最后一个文档
            if current_doc:
                output_file = os.path.join(output_dir, f"section_{file_counter}_{current_heading}.docx")
                current_doc.save(output_file)
                output_files.append(output_file)
            
            logger.info(f"成功分割文档为 {len(output_files)} 个文件")
            return output_files
            
        except Exception as e:
            logger.error(f"分割文档失败: {e}")
            return []
    
    @staticmethod
    def convert_to_html(file_path: str, output_path: str) -> bool:
        """
        将Word文档转换为HTML
        
        Args:
            file_path: Word文档路径
            output_path: HTML输出路径
            
        Returns:
            bool: 是否转换成功
        """
        try:
            doc = Document(file_path)
            html_content = []
            
            html_content.append('<!DOCTYPE html>')
            html_content.append('<html>')
            html_content.append('<head>')
            html_content.append('<meta charset="UTF-8">')
            html_content.append('<title>转换的文档</title>')
            html_content.append('</head>')
            html_content.append('<body>')
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    # 检查是否为标题
                    if paragraph.style and paragraph.style.name.startswith('Heading'):
                        level = int(paragraph.style.name.split()[-1]) if paragraph.style.name.split()[-1].isdigit() else 1
                        html_content.append(f'<h{level}>{paragraph.text}</h{level}>')
                    else:
                        html_content.append(f'<p>{paragraph.text}</p>')
            
            # 添加表格
            for table in doc.tables:
                html_content.append('<table border="1">')
                for row in table.rows:
                    html_content.append('<tr>')
                    for cell in row.cells:
                        html_content.append(f'<td>{cell.text}</td>')
                    html_content.append('</tr>')
                html_content.append('</table>')
            
            html_content.append('</body>')
            html_content.append('</html>')
            
            # 保存HTML文件
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(html_content))
            
            logger.info(f"成功转换为HTML: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"转换为HTML失败: {e}")
            return False
    
    @staticmethod
    def get_document_statistics(file_path: str) -> Dict[str, Any]:
        """
        获取文档统计信息
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict[str, Any]: 统计信息
        """
        try:
            doc = Document(file_path)
            stats = {
                'paragraphs': len(doc.paragraphs),
                'tables': len(doc.tables),
                'images': len(doc.inline_shapes),
                'sections': len(doc.sections),
                'styles': len(doc.styles),
                'word_count': 0,
                'character_count': 0,
                'line_count': 0,
                'page_breaks': 0
            }
            
            # 统计文本信息
            for paragraph in doc.paragraphs:
                text = paragraph.text
                stats['character_count'] += len(text)
                stats['word_count'] += len(text.split())
                stats['line_count'] += text.count('\n') + 1
                
                # 检查分页符
                if '\f' in text or 'page_break' in str(paragraph._element):
                    stats['page_breaks'] += 1
            
            # 统计表格信息
            for table in doc.tables:
                stats['word_count'] += sum(len(cell.text.split()) for row in table.rows for cell in row.cells)
                stats['character_count'] += sum(len(cell.text) for row in table.rows for cell in row.cells)
            
            return stats
            
        except Exception as e:
            logger.error(f"获取文档统计信息失败: {e}")
            return {}


# 使用示例
if __name__ == "__main__":
    # 配置日志
    logging.basicConfig(level=logging.INFO)
    
    # 验证文档
    validation = DocumentUtils.validate_docx_file("example.docx")
    print(f"文档验证结果: {validation}")
    
    # 获取文档摘要
    summary = DocumentUtils.get_document_summary("example.docx")
    print(f"文档摘要: {summary}")
    
    # 提取文本内容
    text = DocumentUtils.extract_text_content("example.docx")
    print(f"文本内容长度: {len(text)}")
    
    # 获取统计信息
    stats = DocumentUtils.get_document_statistics("example.docx")
    print(f"文档统计: {stats}")